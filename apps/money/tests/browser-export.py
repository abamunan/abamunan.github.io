# Browser test for the Excel / PDF report export (More tab). Run from the repo root:
#   python3 -m http.server 8123 &    then    python3 apps/money/tests/browser-export.py     (needs: pip install openpyxl)
import json, os, datetime
os.makedirs("test-output", exist_ok=True)
import openpyxl
from playwright.sync_api import sync_playwright
errs, fails = [], []
def ok(c, m):
    print(("PASS " if c else "FAIL ") + m)
    if not c: fails.append(m)
seed = {"mt_2026_10_7": {"lunch": 400, "acct_lunch": "cash", "ts_lunch": 1, "self": 5000, "acct_self": "bkash", "ts_self": 1},
        "mt_2026_9_20": {"lunch": 100, "acct_lunch": "cash", "ts_lunch": 1}}
with sync_playwright() as p:
    b = p.chromium.launch(); ctx = b.new_context(viewport={"width": 360, "height": 740}, accept_downloads=True)
    ctx.route("**/*", lambda r: r.continue_() if r.request.url.startswith("http://localhost:8123") else r.abort())
    ctx.add_init_script("""(() => { if (!localStorage.getItem('__seeded')) { const s=%s; for (const k in s) localStorage.setItem(k, JSON.stringify(s[k])); localStorage.setItem('__seeded','1'); } })()""" % json.dumps(seed))
    pg = ctx.new_page(); pg.clock.install(time=datetime.datetime(2026, 10, 8, 12, 0, 0))
    pg.on("pageerror", lambda e: errs.append("pageerror: " + str(e)))
    pg.on("console", lambda m: errs.append("console." + m.type + ": " + m.text) if m.type == "error" and "ERR_FAILED" not in m.text and "Failed to load resource" not in m.text else None)
    pg.goto("http://localhost:8123/apps/money/index.html#more"); pg.wait_for_selector("#splash", state="detached", timeout=10000)
    chip = lambda name: pg.locator(".ui-sheet").last.get_by_role("button", name=name, exact=True).click()
    pg.get_by_role("button", name="Export report — Excel / PDF").click(); pg.wait_for_selector(".add-save")
    ok(pg.get_by_role("button", name="2026", exact=True).get_attribute("aria-pressed") == "true" and pg.get_by_role("button", name="Oct", exact=True).get_attribute("aria-pressed") == "true", "defaults to this year and month")
    pg.screenshot(path="test-output/export-0.png")
    # Excel, one month
    with pg.expect_download() as dl: pg.locator(".add-save").click()
    d = dl.value; path = "test-output/" + d.suggested_filename; d.save_as(path)
    ok(d.suggested_filename == "MoneyTracker_October_2026.xlsx", "file name " + d.suggested_filename)
    ws = openpyxl.load_workbook(path).active; vals = [[c for c in r if c is not None] for r in ws.iter_rows(values_only=True)]
    ok(["Total Expense", 400] in vals and ["Total Income", 5000] in vals and ["Net Balance", 4600] in vals, "Excel overview matches the app (400 / 5000 / 4600)")
    ok(["Day 7", 400, 5000, 4600] in vals and ["Lunch", "Expense", 400] in vals, "Excel day row and category row are right")
    ok(ws["A1"].fill.fgColor.rgb == "FF0F3D2E" and ws["B5"].number_format == "#,##0.00", "styles survive (banner fill, money format)")
    pg.wait_for_selector(".add-save", state="detached", timeout=3000)
    # Excel, whole year
    pg.get_by_role("button", name="Export report — Excel / PDF").click(); pg.wait_for_selector(".add-save"); chip("Whole year")
    with pg.expect_download() as dl: pg.locator(".add-save").click()
    d = dl.value; path = "test-output/" + d.suggested_filename; d.save_as(path); ws = openpyxl.load_workbook(path).active; vals = [[c for c in r if c is not None] for r in ws.iter_rows(values_only=True)]
    ok(d.suggested_filename == "MoneyTracker_2026.xlsx" and ["Total Expense", 500] in vals and ["Sep", 100, 0, -100] in vals and ["Oct", 400, 5000, 4600] in vals, "year report: 12 months, Sep and Oct rows, total 500")
    # month with nothing saved → message, sheet stays open, no download
    pg.get_by_role("button", name="Export report — Excel / PDF").click(); pg.wait_for_selector(".add-save"); chip("Jan"); pg.locator(".add-save").click(); pg.wait_for_timeout(400)
    ok(pg.locator(".add-save").count() == 1 and "Nothing saved for January 2026" in pg.inner_text("body"), "empty month: clear message, nothing exported")
    # PDF: printable page in a new tab
    pg.evaluate("window.__printed = 0; const o = window.open; window.open = function () { const w = o.apply(this, arguments); if (w) w.print = () => { window.__printed++; }; return w; }")
    chip("Oct"); chip("PDF (print)")
    with ctx.expect_page() as pop: pg.locator(".add-save").click()
    w = pop.value; w.wait_for_load_state(); txt = w.inner_text("body")
    ok("October 2026" in txt and "400" in txt and "5,000" in txt and "4,600" in txt and "Day 7" in txt, "PDF page shows title, totals and day row")
    pg.wait_for_timeout(700); ok(pg.evaluate("window.__printed") == 1, "print dialog is requested once")
    pdf = w.pdf(format="A4", landscape=True); open("test-output/report.pdf", "wb").write(pdf); ok(pdf[:4] == b"%PDF" and len(pdf) > 2000, "page prints to a real PDF (%d bytes)" % len(pdf))
    ok(not errs, "no console errors: " + "; ".join(errs[:3]))
    b.close()
print("\n%d failed" % len(fails) if fails else "\nall passed"); raise SystemExit(1 if fails else 0)
